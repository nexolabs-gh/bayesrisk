"""Oráculos del contenido final: fuente Study, documento y libros exportados."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from types import SimpleNamespace

from .contract import dump


def expected_artifacts(sc, output: Path) -> dict:
    """Captura el contenido esperado ANTES de leer el archivo que se verificará."""
    from bayesrisk.guided.export import (
        _hojas_de_decisiones,
        _hojas_de_la_etapa,
        _tablas_del_informe,
        decisions_book,
        stage_books,
    )
    from bayesrisk.report.exports import exportable_table
    from bayesrisk.report.renderer import build_document_view

    bundle = sc.study.artifacts.get("report", "input_bundle")
    config = sc.config.report
    view = build_document_view(
        bundle,
        config=config.model_copy(
            update={"html": config.html.model_copy(update={"render_charts": False})}
        ),
    )
    tables = [t for section in view["sections"] for t in section.get("tables", [])]
    report_tables = _tablas_del_informe(sc.study, sc.config.report)
    books = {}
    for stage, name in stage_books().items():
        sheets, indexed = _hojas_de_la_etapa(stage, sc.summary(stage), report_tables)
        books[name] = _sheet_rows(sheets, indexed, exportable_table)
    decisions = _hojas_de_decisiones(sc.project_dir / "run" / "audit_trail.jsonl")
    books[decisions_book()] = _sheet_rows(
        decisions, {k: False for k in decisions}, exportable_table
    )
    expected = {
        "document_tables": tables,
        "document": view,
        "books": books,
        "summaries": {
            stage: {"text": sc.summary(stage).text(), "html": sc.summary(stage)._repr_html_()}
            for stage in sc.steps
        },
    }
    dump(output, expected)
    return expected


def _sheet_rows(sheets, indexed, convert) -> dict:
    from pandas.io.excel._base import ExcelWriter
    from pandas.io.formats.excel import ExcelFormatter

    result = {}
    writer = SimpleNamespace(_date_format="YYYY-MM-DD", _datetime_format="YYYY-MM-DD HH:MM:SS")
    for name, frame in sheets.items():
        table = convert(frame)
        cells = list(ExcelFormatter(table, index=indexed[name]).get_formatted_cells())
        width = max(c.col for c in cells) + 1
        height = max(c.row for c in cells) + 1
        rows = [[None] * width for _ in range(height)]
        for cell in cells:
            value, _ = ExcelWriter._value_with_fmt(writer, cell.val)
            if not isinstance(value, (str, int, float, bool)):
                raise ValueError("tipo de celda Excel no previsto por este corpus")
            rows[cell.row][cell.col] = None if value == "" else value
        result[name] = rows
    return result


def verify_excel(directory: Path, expected: dict) -> int:
    """Censo exacto y cotejo celda a celda; tolerancia sólo de serialización XLSX."""
    import openpyxl
    from openpyxl.compat import safe_string

    if {p.name for p in directory.glob("*.xlsx")} != set(expected):
        raise ValueError("censo de libros Excel diferente")
    cells = 0
    for name, sheets in expected.items():
        book = openpyxl.load_workbook(directory / name, read_only=False, data_only=False)
        try:
            if book.sheetnames != list(sheets):
                raise ValueError(f"censo/orden de hojas diferente: {name}")
            for title, rows in sheets.items():
                if book[title].merged_cells.ranges:
                    raise ValueError("celdas combinadas no previstas por este corpus")
                actual = list(book[title].iter_rows())
                if len(actual) != len(rows):
                    raise ValueError(f"número de filas Excel diferente: {name}/{title}")
                for i, (got, wanted) in enumerate(zip(actual, rows, strict=True)):
                    if len(got) != len(wanted):
                        raise ValueError(f"número de columnas Excel diferente: {name}/{title}")
                    for cell, value in zip(got, wanted, strict=True):
                        if isinstance(value, (list, dict)):
                            value = str(value)
                        if cell.data_type == "f":
                            raise ValueError("fórmula activa inesperada")
                        observed = cell.value
                        equal = (
                            safe_string(observed) == safe_string(value)
                            if isinstance(observed, (float, int))
                            and isinstance(value, (float, int))
                            else observed == (None if value == "" else value)
                        )
                        if not equal:
                            raise ValueError(
                                f"celda Excel distinta: {name}/{title}/{i}/{cell.column}: "
                                f"{observed!r} != {value!r}"
                            )
                        cells += 1
        finally:
            book.close()
    return cells


def _text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _word_tables(document: dict) -> list[dict]:
    from bayesrisk.report.cifras import conteo

    result = []

    def add(key, rows, header=None, caption=None, after=None):
        values = list(rows) if header is None else [header, *rows[:500]]
        result.append(
            {
                "key": key,
                "rows": [[_text(str(c)) for c in row] for row in values],
                "caption": caption,
                "after": after,
            }
        )

    add(
        "cover",
        [(f["label"], f["value"] if f["filled"] else "") for f in document["cover"]]
        + [("Fecha de emisión", document["emitted_date"])],
    )
    add(
        "signatures",
        [("Desarrollo", "", ""), ("Validación independiente", "", ""), ("Aprobación", "", "")],
        ("Rol", "Nombre", "Fecha"),
    )
    for section in document["sections"]:
        if section["kind"] != "summary":
            continue
        summary = section["summary"] or {}
        final = summary.get("final")
        if summary.get("error") or not final:
            continue
        labels = summary["labels"]
        status = [(labels["execution"], final["execution"])]
        if summary.get("assumptions") is None:
            status.append((labels["validation"], final["validation"]))
        add("status", status)
        if final["figures"]:
            add("figures", final["figures"], ("Cifra", "Valor"), labels["figures"])
        add("files", final["files"], caption=labels["files"])
    metrics = document["executive"]["metrics"]
    if metrics:
        add(
            "executive",
            [(m["label"], m["scope"], m["value"], m["band"]) for m in metrics],
            ("Métrica", "Alcance", "Valor", "Estado"),
        )
    for section in document["sections"]:
        if section["kind"] in {"toc", "summary"}:
            continue
        for table in section["tables"]:
            n = min(len(table["rows"]), 500)
            note = (
                f"… (mostrando {conteo(n)} de {table['total_rows_label']} filas)"
                if n < table["total_rows"]
                else None
            )
            add(table["key"], table["rows"], table["columns"], table["title"], note)
        if section["data_exports"] and section["data_exports"]["requested"]:
            raise ValueError("exports de documento no previstos por esta receta")
    return result


def verify_documents(directory: Path, expected: list[dict], document: dict) -> dict:
    """Compara todas las tablas identificadas por dominio, orden y contenido."""
    from lxml import etree, html

    parsed = html.fromstring((directory / "scorecard_report.html").read_text(encoding="utf-8"))
    actual = parsed.xpath("//table[@data-table-key]")
    if [t.get("data-table-key") for t in actual] != [t["key"] for t in expected]:
        raise ValueError("censo/orden de tablas HTML diferente al Study")
    cells = 0
    for node, table in zip(actual, expected, strict=True):
        got = [
            [_text(c.text_content()) for c in row.xpath("./th|./td")] for row in node.xpath(".//tr")
        ]
        wanted = [[_text(str(c)) for c in row] for row in [table["columns"], *table["rows"]]]
        if got != wanted:
            raise ValueError(f"tabla HTML diferente: {table['key']}")
        cells += sum(map(len, wanted))
    structural = []
    for section in document["sections"]:
        summary = section.get("summary") or {}
        final = summary.get("final")
        if final and final["figures"]:
            structural.append(final["figures"])
    if document["executive"]["metrics"]:
        structural.append(
            [
                ("Métrica", "Alcance", "Valor", "Estado"),
                *[
                    (m["label"], m["scope"], m["value"], m["band"])
                    for m in document["executive"]["metrics"]
                ],
            ]
        )
    extra = parsed.xpath("//table[not(@data-table-key)]")
    got_structural = [
        [[_text(c.text_content()) for c in row.xpath("./th|./td")] for row in node.xpath(".//tr")]
        for node in extra
    ]
    want_structural = [[[_text(str(c)) for c in row] for row in table] for table in structural]
    if got_structural != want_structural:
        raise ValueError("tablas estructurales HTML diferentes al Study")
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(directory / "scorecard_report.docx") as archive:
        tree = etree.fromstring(archive.read("word/document.xml"))

    def visible(node):
        pieces = []
        for element in node.iter():
            if element.tag == f"{{{ns['w']}}}t":
                pieces.append(element.text or "")
            elif element.tag in {f"{{{ns['w']}}}br", f"{{{ns['w']}}}tab"}:
                pieces.append(" ")
        return _text("".join(pieces))

    tables = tree.findall("w:body/w:tbl", ns)
    wanted_tables = _word_tables(document)
    if len(tables) != len(wanted_tables):
        raise ValueError("censo de tablas Word diferente al Study")
    word_cells = 0
    for node, table in zip(tables, wanted_tables, strict=True):
        got = [
            [
                _text("\n".join(visible(p) for p in c.findall("w:p", ns)))
                for c in row.findall("w:tc", ns)
            ]
            for row in node.findall("w:tr", ns)
        ]
        if got != table["rows"]:
            raise ValueError(f"tabla Word diferente: {table['key']}")
        for label, neighbor in (("caption", node.getprevious()), ("after", node.getnext())):
            if table[label] and (neighbor is None or visible(neighbor) != _text(table[label])):
                raise ValueError(f"caption/truncado Word distinto: {table['key']}")
        word_cells += sum(map(len, got))
    return {
        "tables": len(expected),
        "word_tables": len(tables),
        "html_cells": cells,
        "word_cells": word_cells,
    }


def verify_artifacts(project: Path, expected: dict) -> dict:
    """El archivo final se verifica contra contenido capturado del Study que lo produjo."""
    return {
        **verify_documents(project / "reports", expected["document_tables"], expected["document"]),
        "excel_cells": verify_excel(project / "excel", expected["books"]),
    }

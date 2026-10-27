"""Reproducciones de la revisión B: magnitud visible, texto literal y extras opcionales."""

import importlib.util
import os
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path

import pandas as pd
import pytest
from test_documentary_layout import bundle_documental

from bayesrisk.report.builder import ReportBuilder
from bayesrisk.report.docx import DocxReportRenderer
from bayesrisk.report.exports import write_workbook
from bayesrisk.report.markdown import MarkdownReportRenderer
from bayesrisk.report.renderer import HtmlReportRenderer


def _bundle_with_binning(variable="edad"):
    bundle, config, _ = bundle_documental(2)
    table = pd.DataFrame({"Bin": ["[0, 1)"], "Event rate": [0.2], "WoE": [0.5], "Count": [10]})
    config = config.model_copy(
        update={"html": config.html.model_copy(update={"render_charts": True})}
    )
    bundle = bundle.model_copy(
        update={
            "cards": {**bundle.cards, "binning": {"summary": {}}},
            "tables": {**bundle.tables, f"binning.tables.{variable}": table},
        }
    )
    return bundle.model_copy(
        update={"sections": ReportBuilder(config).build_sections(bundle)}
    ), config


def test_excel_magnitud_no_nula_visible_en_pvalor_y_soporte(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    values = [0.0, 3.05666320588627e-190, 1e-28, -1e-42, 0.125]
    table = pd.DataFrame({"p-valor": values, "Original": values})
    table.attrs.update(_documentary=True, _documentary_formats={"p-valor": "pvalor"})
    book = openpyxl.load_workbook(
        write_workbook({"Evidencia": table}, tmp_path / "p.xlsx", index=False)
    )
    sheet = book["Evidencia"]
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            assert cell.value == values[cell.row - 2] and cell.data_type == "n"
            if 0 < abs(cell.value) < 1e-15:
                # Un decimal finito no representa estas magnitudes: Excel necesita exponente.
                assert "E+" in cell.number_format.upper(), (cell.coordinate, cell.number_format)
            elif cell.value == 0:
                assert "E" not in cell.number_format.upper()
    book.close()


@pytest.mark.skipif(importlib.util.find_spec("matplotlib") is None, reason="requiere report")
def test_qmd_nombre_de_variable_no_agrega_imagen_externa(tmp_path):
    markdown = pytest.importorskip("markdown")
    variable = "x](https://example.invalid/extra.png) ![resto"
    bundle, config = _bundle_with_binning(variable)
    renderer = MarkdownReportRenderer(config)
    source = renderer.render(bundle)
    renderer.write(source, output_dir=str(tmp_path))
    images = []

    class ImageParser(HTMLParser):
        def handle_starttag(self, tag, attrs):
            if tag == "img":
                images.append(dict(attrs)["src"])

    ImageParser().feed(markdown.markdown(source))
    expected = {p.relative_to(tmp_path).as_posix() for p in tmp_path.glob("*_figuras/*.svg")}
    assert len(expected) >= 2
    assert len(images) == len(expected) and set(images) == expected, images


def test_sin_matplotlib_un_proceso_fresco_conserva_informes_y_tablas():
    code = """
import importlib.abc
import sys
from test_documentary_regressions import _bundle_with_binning
assert 'matplotlib' not in sys.modules
class NoMatplotlib(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname == 'matplotlib' or fullname.startswith('matplotlib.'):
            raise ModuleNotFoundError('matplotlib ausente para este control')
sys.meta_path.insert(0, NoMatplotlib())
from bayesrisk.report.renderer import HtmlReportRenderer
from bayesrisk.report.markdown import MarkdownReportRenderer
bundle, config = _bundle_with_binning()
assert '[0, 1)' in HtmlReportRenderer(config).render(bundle)
assert '[0, 1)' in MarkdownReportRenderer(config).render(bundle).replace(chr(92), '')
try:
    import docx
except ImportError:
    pass
else:
    from bayesrisk.report.docx import DocxReportRenderer
    assert DocxReportRenderer(config).render(bundle).startswith(b'PK')
assert 'matplotlib' not in sys.modules
"""
    env = {**os.environ, "PYTHONPATH": os.pathsep.join([str(Path(__file__).parent), *sys.path])}
    result = subprocess.run([sys.executable, "-c", code], env=env, capture_output=True, text=True)
    assert result.returncode == 0, result.stdout + result.stderr


def test_error_de_un_grafico_no_elimina_las_tablas(monkeypatch):
    from bayesrisk.report import charts

    def broken(*args, **kwargs):
        raise RuntimeError("gráfico indisponible")

    monkeypatch.setattr(charts, "_binning_panel", broken)
    bundle, config = _bundle_with_binning()
    assert "[0, 1)" in HtmlReportRenderer(config).render(bundle)
    assert "[0, 1)" in MarkdownReportRenderer(config).render(bundle).replace("\\", "")
    if importlib.util.find_spec("docx") is not None:
        assert DocxReportRenderer(config).render(bundle).startswith(b"PK")

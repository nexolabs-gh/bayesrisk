"""D-INF-1…4 (`docs/design/_ENMIENDA-INFORME-LEGIBLE.md`): el informe se lee en es-CL.

Cada número que el informe imprime —tablas, listas, anexo de parámetros, linaje, gráficos y la
página ejecutiva— pasa por :mod:`nikodym.report.cifras`: coma decimal, cuatro decimales, miles con
punto en los conteos, «sí»/«no», y la regla del cero final, que impide que un redondeo se haga pasar
por un corte (un PSI ``0.24996`` no se escribe ``0,2500`` junto a su corte de ``0,25``).
"""

from __future__ import annotations

import io
import re
from decimal import Decimal
from typing import Any

import pandas as pd
import pytest
from test_report_charts import _stability_frame
from test_report_renderer import _bundle, _renderer

from nikodym.report import charts, prose
from nikodym.report import renderer as renderer_module
from nikodym.report.cifras import (
    cifra,
    conteo,
    corte,
    es_columna_de_conteo,
    es_columna_de_pvalor,
    pvalor,
)

# ───────────────────────────── la regla ─────────────────────────────


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        # Base: cuatro decimales entre 0,001 y 1.000.
        (0.042677, "0,0427"),
        (452.337071, "452,3371"),
        (-1.163851, "-1,1639"),
        # Desde 1.000: dos decimales y miles agrupados.
        (697376973.922913, "697.376.973,92"),
        (1802.42, "1.802,42"),
        # Bajo 0,001: dos cifras significativas, nunca «0,0000» para algo que no es cero.
        (0.000015, "0,000015"),
        (0.00049, "0,00049"),
        # Bajo una millonésima: científica con coma.
        (2.3e-09, "2,3e-09"),
        (3e-13, "3,0e-13"),
        (9.96e-07, "9,96e-07"),  # la mantisa no sube a «10,0e-07»: la extiende el cero final
        # Cero, también -0.0, sin signo.
        (0.0, "0,0000"),
        (-0.0, "0,0000"),
        # Ausentes e infinitos.
        (None, "—"),
        (float("nan"), "—"),
        (float("inf"), "inf"),
        (float("-inf"), "-inf"),
    ],
)
def test_la_base_de_la_regla(valor: float | None, esperado: str) -> None:
    assert cifra(valor) == esperado


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        # Un redondeo que termina en cero y no es exacto se extiende: se leería «en el corte».
        (0.24996, "0,24996"),
        (0.25004, "0,25004"),
        (0.05004, "0,05004"),
        (999.9999999, "999,9999999"),
        # Un exacto no se toca, aunque termine en cero.
        (0.25, "0,2500"),
        (0.1, "0,1000"),  # el exacto es `repr` = 0.1, no el binario 0.1000000000000000055…
        (1000.0, "1.000,00"),
    ],
)
def test_la_regla_del_cero_final(valor: float, esperado: str) -> None:
    assert cifra(valor) == esperado


def test_un_decimal_no_pasa_por_float() -> None:
    """En float, ``Decimal("0.24999999999999999999")`` es 0.25 y la regla lo daría por exacto."""
    assert cifra(Decimal("0.24999999999999999999")) == "0,24999999999999999999"
    assert cifra(Decimal("1E-400")) == "1,0e-400"  # en float es cero
    assert cifra(Decimal("4338485154.07")) == "4.338.485.154,07"
    assert cifra(Decimal("NaN")) == "—"


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        (0.000096, "< 0,001"),
        (0.0342, "0,034"),
        (0.04996, "0,04996"),  # junto al corte rojo 0,05: no se escribe «0,050»
        (0.05, "0,050"),  # exacto
        (None, "—"),
        (float("nan"), "—"),
    ],
)
def test_los_p_valores(valor: float | None, esperado: str) -> None:
    assert pvalor(valor) == esperado


def test_los_cortes_del_config_van_exactos() -> None:
    assert corte(0.125) == "0,125"
    assert corte(0.1) == "0,10"
    assert corte(0.25) == "0,25"
    assert corte(0.05004) == "0,05004"
    assert corte(5.0, minimo=1) == "5,0"
    assert corte(Decimal("0.24999999999999999999")) == "0,24999999999999999999"


def test_que_columnas_cuentan_y_cuales_son_p_valores() -> None:
    """La lista de conteos va en la dirección segura: un identificador nunca se agrupa."""
    for nombre in (
        "n",
        "n_total",
        "N_BAD",
        "cum_good",
        "Count",
        "expected_count",
        "n_rows",
        "Event",
    ):
        assert es_columna_de_conteo(nombre), nombre
    for nombre in ("tramo", "period", "cohort", "stage", "root_seed", "year", "loan_id", "decile"):
        assert not es_columna_de_conteo(nombre), nombre
    assert conteo(30316) == "30.316"
    assert es_columna_de_pvalor("p_value") and es_columna_de_pvalor("wald_pvalue")
    assert not es_columna_de_pvalor("value")


# ───────────────────────────── las tablas ─────────────────────────────


def test_una_tabla_se_escribe_en_es_cl_y_marca_sus_columnas_numericas() -> None:
    tabla = pd.DataFrame(
        {
            "tramo": [2005, 2006],
            "n_total": [30316, 7733],
            "default_rate": [0.2380, 0.24996],
            "p_value": [0.000096, 0.0342],
            "selected": [True, False],
            "name": ["plazo_meses", "sector"],
        }
    )
    vista = renderer_module._table_view("eda.univariate.profiles.anio_fiscal", tabla, max_rows=10)
    assert vista["rows"] == [
        ("2005", "30.316", "0,2380", "< 0,001", "sí", "plazo_meses"),
        ("2006", "7.733", "0,24996", "0,034", "no", "sector"),
    ]
    # Numéricas: el año, el conteo, la tasa y el p-valor; no el booleano ni el texto.
    assert vista["numeric"] == [True, True, True, True, False, False]


def test_el_linaje_no_agrupa_la_semilla_y_el_json_conserva_su_sintaxis() -> None:
    assert renderer_module._display_scalar(20260920, key_path=("root_seed",)) == "20260920"
    assert renderer_module._display_json_value(
        {"alpha": 0.05004, "enabled": True, "n_bins": 5}, key_path=("model",)
    ) == {"alpha": "0,05004", "enabled": True, "n_bins": 5}


def test_el_informe_no_tiene_celdas_con_punto_decimal_ni_true_false() -> None:
    """El HTML del bundle sintético completo, contado celda por celda.

    El extractor se prueba aparte (abajo) para que este test no nazca verde por un regex que no
    encuentra celdas.
    """
    html = _renderer().render(_bundle())
    celdas = _celdas(html)
    assert len(celdas) > 30  # 36 en el bundle sintético
    assert [c for c in celdas if re.fullmatch(r"-?\d+\.\d+", c)] == []
    assert [c for c in celdas if c in {"true", "false"}] == []
    assert 'class="num"' in html


def _celdas(html: str) -> list[str]:
    return [c.strip() for c in re.findall(r"<td(?:\s[^>]*)?>(.*?)</td>", html, re.S)]


def test_el_extractor_de_celdas_encuentra_las_de_punto_decimal() -> None:
    muestra = '<tr><td>0.042677</td><td class="num">true</td><td>0,0427</td></tr>'
    assert _celdas(muestra) == ["0.042677", "true", "0,0427"]


def test_el_word_alinea_a_la_derecha_las_columnas_numericas() -> None:
    docx = pytest.importorskip("docx")
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    from nikodym.report.docx import DocxReportRenderer

    documento = docx.Document(io.BytesIO(DocxReportRenderer().render(_bundle())))
    derecha: set[str] = set()
    izquierda: set[str] = set()
    for tabla in documento.tables:
        encabezados = [celda.text for celda in tabla.rows[0].cells]
        for indice, encabezado in enumerate(encabezados):
            alineaciones = {
                fila.cells[indice].paragraphs[0].alignment for fila in list(tabla.rows)[1:]
            }
            if alineaciones == {WD_ALIGN_PARAGRAPH.RIGHT}:
                derecha.add(encabezado)
            elif WD_ALIGN_PARAGRAPH.RIGHT not in alineaciones:
                izquierda.add(encabezado)
    assert "beta" in derecha and "p_value" in derecha
    assert "feature" in izquierda


# ───────────────────────────── la página ejecutiva ─────────────────────────────


def _con_estabilidad(stable: float, review: float, psi: float) -> Any:
    bundle = _bundle()
    cards = dict(bundle.cards)
    cards["stability"] = {
        "comparisons": ["dev_vs_oot"],
        "max_psi_by_comparison": {"dev_vs_oot": psi},
        "psi_metric_by_comparison": {"dev_vs_oot": "score_psi"},
        "bands_by_comparison": {"dev_vs_oot": "review"},
        "stable_threshold": stable,
        "review_threshold": review,
    }
    return bundle.model_copy(update={"cards": cards})


def test_la_pagina_ejecutiva_no_hace_pasar_un_psi_por_su_corte() -> None:
    vista = prose.executive_view(_con_estabilidad(0.125, 0.25, 0.24996))
    psi = [m for m in vista.metrics if m.label.startswith("Peor PSI")]
    assert [m.value for m in psi] == ["0,24996"]
    nota = next(n for n in vista.notes if "bandas de PSI" in n)
    assert "estable por debajo de 0,125" in nota and "redesarrollo desde 0,25" in nota


def test_con_los_umbrales_por_defecto_la_nota_no_cambia() -> None:
    vista = prose.executive_view(_con_estabilidad(0.10, 0.25, 0.1612))
    nota = next(n for n in vista.notes if "bandas de PSI" in n)
    assert "estable por debajo de 0,10, revisión desde 0,10 y redesarrollo desde 0,25." in nota
    assert [m.value for m in vista.metrics if m.label.startswith("Peor PSI")] == ["0,1612"]


# ───────────────────────────── los gráficos ─────────────────────────────


def test_la_leyenda_de_estabilidad_escribe_el_corte_exacto() -> None:
    frame = _stability_frame().assign(stable_threshold=0.125)
    svg = charts.render_stability_chart(frame, title="Estabilidad")
    assert "Revisión: 0,125 ≤ índice &lt; 0,25" in svg
    assert "0,12 ≤" not in svg


def test_brier_y_ece_no_se_leen_cero() -> None:
    assert charts._reliability_label("Desarrollo", {"brier": 0.0004, "ece": 0.0123}) == (
        "Desarrollo (Brier=0,00040, ECE=0,0123)"
    )


@pytest.mark.parametrize(("escala", "sufijo"), [(1e-9, "e-09"), (1e-13, "e-13")])
def test_un_eje_de_coeficientes_diminutos_no_rotula_ceros(escala: float, sufijo: str) -> None:
    """Con dos decimales fijos, un forest de ``3e-09`` rotulaba ``0,00`` todas sus marcas; y un
    redondeo absoluto a doce decimales volvía a hacerlo con ``3e-13`` (pasada 2 de Codex)."""
    coeficientes = [
        {"feature": "monto", "beta": 3 * escala, "conf_low": 2 * escala, "conf_high": 4 * escala},
        {
            "feature": "plazo",
            "beta": 1 * escala,
            "conf_low": 0.5 * escala,
            "conf_high": 1.5 * escala,
        },
    ]
    svg = str(charts.render_coefficients_forest(coeficientes, title="Coeficientes"))
    assert sufijo in svg
    assert ">0,00<" not in svg and "0,00 " not in svg


def test_las_marcas_del_eje_se_limpian_relativo_al_paso() -> None:
    assert charts._marcas_del_eje([0.0, 0.05, 0.1, 0.15000000000000002], 2) == [
        "0,00",
        "0,05",
        "0,10",
        "0,15",
    ]
    assert charts._marcas_del_eje([0.0, 0.125, 0.25], 2) == ["0,000", "0,125", "0,250"]
    assert charts._marcas_del_eje([-0.2, -0.1, 2.7755575615628914e-17, 0.1, 0.2], 2) == [
        "-0,20",
        "-0,10",
        "0,00",
        "0,10",
        "0,20",
    ]
    assert charts._marcas_del_eje([2e-13, 2.5e-13, 3e-13], 2) == ["2,0e-13", "2,5e-13", "3,0e-13"]


# ───────────────── pasada 1 de Codex sobre el código ─────────────────


@pytest.mark.parametrize(("psi", "texto"), [(0.24996, "0,24996"), (0.25004, "0,25004")])
def test_el_cuerpo_y_las_conclusiones_no_hacen_pasar_un_psi_por_su_corte(
    psi: float, texto: str
) -> None:
    """La página ejecutiva ya usaba la regla; el cuerpo de estabilidad y las conclusiones
    escribían el PSI con cuatro decimales fijos junto a su banda («0,2500»)."""
    bundle = _con_estabilidad(0.10, 0.25, psi)
    cuerpo = " ".join(prose._results_stability(bundle))
    conclusiones = " ".join(prose.conclusions_body(bundle))
    en_el_corte = re.compile(r"0,2500(?!\d)")  # «0,2500» y no el prefijo de «0,25004»
    assert texto in cuerpo and not en_el_corte.search(cuerpo)
    assert texto in conclusiones and not en_el_corte.search(conclusiones)


def test_un_decimal_de_mas_de_28_cifras_no_se_redondea_al_corte() -> None:
    """Con el contexto de 28 dígitos, ``normalize`` redondeaba este Decimal a 0.25 y la regla del
    cero final se detenía en «0,2500»; y ``Decimal("1E+24")`` levantaba ``InvalidOperation``."""
    assert cifra(Decimal("0.24999999999999999999999999999")) == "0,24999999999999999999999999999"
    assert cifra(Decimal("-0.24999999999999999999999999999")).startswith("-0,2499")
    assert cifra(Decimal("1E+24")) == "1.000.000.000.000.000.000.000.000,00"


def test_el_aviso_de_truncado_agrupa_sus_conteos() -> None:
    tabla = pd.DataFrame({"n_total": range(30316)})
    vista = renderer_module._table_view("x.frame", tabla, max_rows=1000)
    assert vista["truncated"] is True
    assert (vista["shown_rows_label"], vista["total_rows_label"]) == ("1.000", "30.316")
    assert vista["total_rows"] == 30316  # el entero sigue para la lógica


def test_el_config_del_anexo_c_se_escribe_exacto_y_no_como_p_valor_observado() -> None:
    """Pasada 2 de Codex sobre el código: el config efectivo del Anexo C pasaba por la regla de
    los p-valores observados, y un `entry_p_value` de 0.0005 se leía «< 0,001»."""
    config = {"stepwise": {"entry_p_value": 0.0005, "exit_p_value": 0.1254}, "tol": 1e-08}
    assert renderer_module._display_json_value(config, key_path=("effective_config",)) == {
        "stepwise": {"entry_p_value": "0,0005", "exit_p_value": "0,1254"},
        "tol": "0,00000001",
    }
    # Un p-valor observado de una card sí sigue la regla de D-CPY-4.
    assert renderer_module._display_json_value({"p_value": 0.0005}, key_path=("card",)) == {
        "p_value": "< 0,001"
    }
    for corte_del_config in ("entry_p_value", "exit_p_value", "max_pvalue", "p_value_threshold"):
        assert not es_columna_de_pvalor(corte_del_config), corte_del_config


def test_un_corte_repetido_en_una_card_se_lee_igual_que_en_el_config() -> None:
    """Pasada 3 de Codex sobre el código: los `thresholds` de las cards se redondeaban junto al
    `effective_config` exacto, y el Anexo C mostraba dos valores para el mismo corte."""
    card = renderer_module._display_json_value(
        {"thresholds": {"entry_p_value": 0.12541}, "ks_cutoff_score": 0.210709123},
        key_path=("model",),
    )
    config = renderer_module._display_json_value(
        {"stepwise": {"entry_p_value": 0.12541}}, key_path=("effective_config",)
    )
    assert card["thresholds"]["entry_p_value"] == config["stepwise"]["entry_p_value"] == "0,12541"
    # `ks_cutoff_score` contiene «cut» pero es un resultado: sigue la regla de las cifras.
    assert card["ks_cutoff_score"] == "0,2107"
    # En una tabla, los umbrales se leen como en la prosa.
    assert renderer_module._display_scalar(0.1, key_path=("t", "stable_threshold")) == "0,10"
    assert renderer_module._display_scalar(0.05, key_path=("t", "alpha")) == "0,05"

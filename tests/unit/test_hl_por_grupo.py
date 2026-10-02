"""D-HLG-1…3: Hosmer-Lemeshow dice su magnitud (enmienda HL-Y-DECISIONES-EN-EL-YAML §1.2 (a)).

Tres contratos, del kernel al artefacto que lee la persona:

1. **D-HLG-1** — el kernel publica sus grupos (``n``, malos observados y esperados, tasa
   observada, PD media, diferencia en puntos porcentuales, O/E y contribución al estadístico) con
   **los mismos grupos que usó el test** —orden estable por PD y ``np.array_split``, no ``qcut``—, y
   el paso de validación los publica en la clave aditiva ``("validation",
   "hosmer_lemeshow_groups")``, muestra por grupo, sólo para las muestras con veredicto.
2. **D-HLG-2** — la línea de un Hosmer-Lemeshow que falla suma el grupo de **mayor diferencia
   absoluta** entre la tasa observada y la PD media (a igual diferencia, el de menor número), con
   sus dos tasas; **no** el de mayor contribución al estadístico (en el SBA, el 8 y no el 3). El
   informe (HTML y Word) gana una tabla por muestra en la subsección de calibración y la pantalla
   recibe la clave en su payload.
3. Ningún veredicto cambia: lo comprueba la proyección canónica del preset F1 (sólo la clave
   nueva), fuera de este archivo.

Contrato: ``docs/design/_ENMIENDA-HL-Y-DECISIONES-EN-EL-YAML.md`` §1.2 y §6.
"""

from __future__ import annotations

import importlib.util
import io
import math
import re
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import numpy as np
import pandas as pd
import pytest

from bayesrisk.guided.summaries import _pruebas_decisivas
from bayesrisk.report.builder import ReportBuilder
from bayesrisk.report.config import ReportConfig
from bayesrisk.report.prose import _PARTITION_LABELS
from bayesrisk.report.renderer import HtmlReportRenderer
from bayesrisk.validation import calibration_tests
from bayesrisk.validation import results as validation_results
from bayesrisk.validation.calibration_tests import hosmer_lemeshow
from bayesrisk.validation.config import CalibrationValidationConfig, ValidationConfig
from bayesrisk.validation.evaluator import ValidationEvaluator
from bayesrisk.validation.exceptions import CalibrationTestError
from bayesrisk.validation.step import VALIDATION_ARTIFACTS, ValidationStep

_HAS_DOCX = importlib.util.find_spec("docx") is not None


def hosmer_lemeshow_groups(y: np.ndarray, p: np.ndarray, *, n_groups: int) -> pd.DataFrame:
    # Se resuelve al llamar: así cada test nace rojo por su cuenta sobre el HEAD anterior.
    return calibration_tests.hosmer_lemeshow_groups(y, p, n_groups=n_groups)  # type: ignore[attr-defined]


#: Las columnas del artefacto, en su orden (literal a propósito: es el contrato publicado).
_COLUMNAS: tuple[str, ...] = (
    "partition",
    "group",
    "n",
    "observed_defaults",
    "expected_defaults",
    "observed_dr",
    "mean_pd",
    "gap_pp",
    "oe_ratio",
    "contribution",
)


# ───────────────────────── D-HLG-1: el kernel publica sus grupos ─────────────────────────


def test_las_columnas_publicadas_son_el_contrato() -> None:
    assert tuple(validation_results.HOSMER_LEMESHOW_GROUP_COLUMNS) == _COLUMNAS


def test_los_grupos_son_los_del_test_golden_a_mano() -> None:
    """20 operaciones en 3 grupos de 7, 7 y 6: el grupo 2 parte el bloque de PD 0,2 (``qcut`` no
    puede partir un empate), así que el golden distingue los grupos del test de cualquier corte
    por valor."""
    p = np.repeat([0.05, 0.1, 0.2, 0.3], [7, 3, 6, 4]).astype(float)
    y = np.array([0, 0, 0, 0, 0, 0, 1, 0, 0, 1, 0, 0, 1, 0, 0, 1, 0, 1, 1, 0], dtype=float)
    tabla = hosmer_lemeshow_groups(y, p, n_groups=3)
    assert tuple(tabla.columns) == _COLUMNAS[1:]
    assert tabla["group"].tolist() == [1, 2, 3]
    assert tabla["n"].tolist() == [7, 7, 6]
    assert tabla["observed_defaults"].tolist() == [1, 2, 3]
    esperados = [7 * 0.05, 3 * 0.1 + 4 * 0.2, 2 * 0.2 + 4 * 0.3]
    medias = [0.05, 1.1 / 7, 1.6 / 6]
    assert tabla["expected_defaults"].tolist() == pytest.approx(esperados, rel=1e-12)
    assert tabla["mean_pd"].tolist() == pytest.approx(medias, rel=1e-12)
    assert tabla["observed_dr"].tolist() == pytest.approx([1 / 7, 2 / 7, 3 / 6], rel=1e-12)
    assert tabla["gap_pp"].tolist() == pytest.approx(
        [100 * (1 / 7 - 0.05), 100 * (2 / 7 - 1.1 / 7), 100 * (3 / 6 - 1.6 / 6)], rel=1e-12
    )
    assert tabla["oe_ratio"].tolist() == pytest.approx([1 / 0.35, 2 / 1.1, 3 / 1.6], rel=1e-12)
    contribuciones = [
        (1 - 0.35) ** 2 / (7 * 0.05 * 0.95),
        (2 - 1.1) ** 2 / (7 * (1.1 / 7) * (1 - 1.1 / 7)),
        (3 - 1.6) ** 2 / (6 * (1.6 / 6) * (1 - 1.6 / 6)),
    ]
    assert tabla["contribution"].tolist() == pytest.approx(contribuciones, rel=1e-12)
    registro = hosmer_lemeshow(y, p, n_groups=3)
    assert registro.statistic is not None
    assert math.isclose(sum(contribuciones), registro.statistic, rel_tol=1e-12)


def test_la_suma_de_las_contribuciones_reproduce_el_estadistico() -> None:
    rng = np.random.default_rng(7)
    p = rng.uniform(0.02, 0.6, size=1000)
    y = (rng.uniform(size=1000) < np.clip(p * 1.15, 0, 1)).astype(float)
    registro = hosmer_lemeshow(y, p, n_groups=10)
    tabla = hosmer_lemeshow_groups(y, p, n_groups=10)
    assert registro.statistic is not None
    assert math.isclose(float(tabla["contribution"].sum()), registro.statistic, rel_tol=1e-12)
    assert int(tabla["n"].sum()) == 1000
    assert int(tabla["observed_defaults"].sum()) == int(y.sum())


def test_sin_grupos_validos_no_hay_tabla() -> None:
    with pytest.raises(CalibrationTestError):
        hosmer_lemeshow_groups(np.array([0.0, 1.0]), np.array([0.2, 0.3]), n_groups=5)


def _frame_dos_muestras() -> pd.DataFrame:
    rng = np.random.default_rng(11)
    desarrollo_p = rng.uniform(0.03, 0.5, size=300)
    desarrollo_y = (rng.uniform(size=300) < desarrollo_p).astype(int)
    return pd.DataFrame(
        {
            "partition": ["desarrollo"] * 300 + ["oot"] * 8,
            "target": [*desarrollo_y.tolist(), 0, 1, 0, 0, 1, 0, 0, 0],
            "pd_calibrated": [*desarrollo_p.tolist(), *([0.2] * 8)],
        },
        index=[f"r{i}" for i in range(308)],
    )


def _evaluador(**calibracion: Any) -> ValidationEvaluator:
    cfg = ValidationConfig(
        families=("calibration",),
        calibration=CalibrationValidationConfig(
            hl_n_groups=5, min_rows_per_group=10, binomial_by_grade=False, **calibracion
        ),
    )
    return ValidationEvaluator.from_config(cfg)


def test_el_evaluador_publica_solo_las_muestras_con_veredicto() -> None:
    """La OOT de 8 filas queda sin veredicto (bajo el mínimo): no tiene grupos que mostrar."""
    evaluador = _evaluador()
    frame = _frame_dos_muestras()
    resultado = evaluador.validate(calibrated_pd=frame)
    tabla = evaluador.hosmer_lemeshow_groups(frame, resultado.calibration_records)
    assert tuple(tabla.columns) == _COLUMNAS
    assert set(tabla["partition"]) == {"desarrollo"}
    assert tabla["group"].tolist() == [1, 2, 3, 4, 5]
    (hl,) = [
        r
        for r in resultado.calibration_records
        if r.test == "hosmer_lemeshow" and r.partition == "desarrollo"
    ]
    assert hl.statistic is not None
    assert math.isclose(float(tabla["contribution"].sum()), hl.statistic, rel_tol=1e-12)


def test_sin_hosmer_lemeshow_la_tabla_sale_vacia_con_sus_columnas() -> None:
    evaluador = _evaluador(hosmer_lemeshow=False)
    frame = _frame_dos_muestras()
    resultado = evaluador.validate(calibrated_pd=frame)
    tabla = evaluador.hosmer_lemeshow_groups(frame, resultado.calibration_records)
    assert tabla.empty
    assert tuple(tabla.columns) == _COLUMNAS


def test_la_clave_es_aditiva_y_la_declara_el_paso() -> None:
    assert "hosmer_lemeshow_groups" in VALIDATION_ARTIFACTS
    assert ("validation", "hosmer_lemeshow_groups") in ValidationStep.provides
    # Las claves de siempre siguen ahí, en su orden.
    assert VALIDATION_ARTIFACTS[:6] == (
        "discrimination",
        "calibration",
        "stability",
        "backtesting",
        "result",
        "card",
    )


# ───────────────────────── D-HLG-2: la frase dice la mayor diferencia ─────────────────────────

#: Los diez grupos de Desarrollo del SBA de Cami (medido el 2026-09-29 y el 2026-10-02): la mayor
#: diferencia absoluta es el grupo 8 (3,42 pp) y la mayor contribución, el grupo 3 (O/E 0,64).
_SBA_DESARROLLO: tuple[tuple[int, int, int, float, float, float], ...] = (
    # grupo, n, malos, esperados, tasa observada, PD media
    (1, 3032, 43, 44.4, 0.014182, 0.014644),
    (2, 3032, 69, 96.4, 0.022757, 0.031794),
    (3, 3032, 89, 140.2, 0.029354, 0.046240),
    (4, 3032, 197, 196.3, 0.064974, 0.064743),
    (5, 3032, 274, 265.4, 0.090369, 0.087533),
    (6, 3032, 379, 383.2, 0.125000, 0.126385),
    (7, 3031, 751, 709.9, 0.247773, 0.234213),
    (8, 3031, 1381, 1277.4, 0.455625, 0.421445),
    (9, 3031, 1804, 1788.9, 0.595183, 0.590201),
    (10, 3031, 2227, 2311.8, 0.734741, 0.762719),
)


def _grupos(filas: tuple[tuple[int, int, int, float, float, float], ...]) -> pd.DataFrame:
    registros = []
    for grupo, n, malos, esperados, tasa, media in filas:
        contribucion = (malos - esperados) ** 2 / (n * media * (1 - media))
        registros.append(
            {
                "partition": "desarrollo",
                "group": grupo,
                "n": n,
                "observed_defaults": malos,
                "expected_defaults": esperados,
                "observed_dr": tasa,
                "mean_pd": media,
                "gap_pp": 100 * (tasa - media),
                "oe_ratio": malos / esperados,
                "contribution": contribucion,
            }
        )
    return pd.DataFrame(registros, columns=list(_COLUMNAS))


class _Almacen:
    def __init__(self, valores: dict[tuple[str, str], Any]) -> None:
        self._valores = valores

    def has(self, domain: str, key: str) -> bool:
        return (domain, key) in self._valores

    def get(self, domain: str, key: str) -> Any:
        return self._valores[(domain, key)]


def _study(valores: dict[tuple[str, str], Any]) -> Any:
    return SimpleNamespace(artifacts=_Almacen(valores), config=SimpleNamespace())


def _calibracion_que_falla() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "test": "hosmer_lemeshow",
                "partition": "desarrollo",
                "decision": "fail",
                "p_value": 6.9e-10,
                "expected_pd": 0.237967,
                "observed_dr": 0.237960,
            }
        ]
    )


def test_la_frase_nombra_el_grupo_de_mayor_diferencia_absoluta_no_el_de_mayor_contribucion() -> (
    None
):
    grupos = _grupos(_SBA_DESARROLLO)
    # Premisa del caso (pasada 1 de Codex sobre la enmienda): los dos criterios difieren.
    assert int(grupos.loc[grupos["contribution"].idxmax(), "group"]) == 3
    assert int(grupos.loc[grupos["gap_pp"].abs().idxmax(), "group"]) == 8
    (linea,) = _pruebas_decisivas(
        _study(
            {
                ("validation", "calibration"): _calibracion_que_falla(),
                ("validation", "hosmer_lemeshow_groups"): grupos,
            }
        )
    )
    assert linea == (
        "Hosmer-Lemeshow en Desarrollo (p-valor < 0,001; PD media agregada 23,80 % frente a "
        "23,80 % observada; la mayor diferencia, en el grupo 8 de 10: 45,56 % observado frente "
        "a 42,14 % predicho)"
    )


def test_a_igual_diferencia_gana_el_grupo_de_menor_numero() -> None:
    filas = (
        (1, 100, 10, 8.0, 0.10, 0.08),
        (2, 100, 20, 22.0, 0.20, 0.22),
        (3, 100, 40, 40.0, 0.40, 0.40),
    )
    (linea,) = _pruebas_decisivas(
        _study(
            {
                ("validation", "calibration"): _calibracion_que_falla(),
                ("validation", "hosmer_lemeshow_groups"): _grupos(filas),
            }
        )
    )
    assert (
        "la mayor diferencia, en el grupo 1 de 3: 10,00 % observado frente a 8,00 % predicho"
        in linea
    )


def test_un_estudio_sin_la_clave_conserva_la_frase_de_antes() -> None:
    """Una corrida guardada antes de la clave no trae los grupos: la frase es la de D-CPY-6."""
    (linea,) = _pruebas_decisivas(_study({("validation", "calibration"): _calibracion_que_falla()}))
    assert linea == (
        "Hosmer-Lemeshow en Desarrollo (p-valor < 0,001; PD media agregada 23,80 % frente a "
        "23,80 % observada)"
    )


# ───────────────────────── D-HLG-2: el informe y la pantalla, sobre una corrida real ──────────


@pytest.fixture(autouse=True)
def _usar_fake_binning_process(fake_binning_process: object) -> None:
    del fake_binning_process


@pytest.fixture(scope="module")
def corrida(tmp_path_factory: pytest.TempPathFactory) -> Any:
    from test_report_pagina_ejecutiva import _corrida

    return _corrida(tmp_path_factory.mktemp("hl_por_grupo"), name="hl", decidir=False)


def _muestras_con_veredicto(study: Any) -> list[str]:
    calibracion = study.artifacts.get("validation", "calibration")
    hl = calibracion[(calibracion["test"] == "hosmer_lemeshow") & calibracion["statistic"].notna()]
    return [str(p) for p in hl["partition"]]


def test_la_corrida_publica_los_grupos_de_cada_muestra_con_veredicto(corrida: Any) -> None:
    study = corrida.study
    tabla = study.artifacts.get("validation", "hosmer_lemeshow_groups")
    muestras = _muestras_con_veredicto(study)
    assert muestras, "la corrida de prueba debe tener al menos un Hosmer-Lemeshow con veredicto"
    assert list(dict.fromkeys(tabla["partition"])) == muestras
    calibracion = study.artifacts.get("validation", "calibration")
    for muestra in muestras:
        (estadistico,) = calibracion[
            (calibracion["test"] == "hosmer_lemeshow") & (calibracion["partition"] == muestra)
        ]["statistic"]
        suma = float(tabla.loc[tabla["partition"] == muestra, "contribution"].sum())
        assert math.isclose(suma, float(estadistico), rel_tol=1e-12)


def _tabla_html(html: str, key: str) -> str | None:
    patron = rf'<table[^>]*data-table-key="{re.escape(key)}"[^>]*>(.*?)</table>'
    encontrada = re.search(patron, html, re.S)
    return None if encontrada is None else encontrada.group(1)


def test_el_informe_html_trae_una_tabla_por_muestra_en_calibracion(corrida: Any) -> None:
    cfg = ReportConfig(sections={"missing_policy": "skip"})
    bundle = ReportBuilder(cfg).collect(corrida.study)
    html = HtmlReportRenderer(cfg).render(bundle)
    # La subsección de calibración, de su apertura a la sección siguiente: ahí tienen que estar.
    inicio = html.find('data-section-id="validation.calibration"')
    assert inicio >= 0
    fin = html.find("<section ", inicio)
    subseccion = html[inicio : fin if fin >= 0 else len(html)]
    assert 'data-table-key="validation.calibration"' in subseccion
    for muestra in _muestras_con_veredicto(corrida.study):
        clave = f"validation.hosmer_lemeshow_groups.{muestra}"
        tabla = _tabla_html(html, clave)
        assert tabla is not None, clave
        encabezados = re.findall(r"<th(?:\s[^>]*)?>(.*?)</th>", tabla)
        assert encabezados[:2] == ["Grupo", "Observaciones"]
        assert "Diferencia (pp)" in encabezados and "O/E" in encabezados
        # En el cuerpo, junto a la tabla de calibración, y no en el anexo de tablas.
        assert f'data-table-key="{clave}"' in subseccion
        assert html.count(f'data-table-key="{clave}"') == 1


@pytest.mark.skipif(not _HAS_DOCX, reason="requiere python-docx (extra docx)")
def test_el_informe_word_trae_las_mismas_tablas(corrida: Any) -> None:
    import docx

    from bayesrisk.report.docx import DocxReportRenderer

    cfg = ReportConfig(sections={"missing_policy": "skip"})
    bundle = ReportBuilder(cfg).collect(corrida.study)
    documento = docx.Document(io.BytesIO(DocxReportRenderer(cfg).render(bundle)))
    texto = "\n".join(p.text for p in documento.paragraphs)
    for muestra in _muestras_con_veredicto(corrida.study):
        assert f"Hosmer-Lemeshow por grupo · {_PARTITION_LABELS[muestra]}" in texto
    encabezados = {
        tuple(c.text for c in tabla.rows[0].cells) for tabla in documento.tables if tabla.rows
    }
    assert any(fila[:2] == ("Grupo", "Observaciones") for fila in encabezados)


def test_la_pantalla_recibe_los_grupos(corrida: Any) -> None:
    from bayesrisk.ui.serializers import serialize_study

    payload = serialize_study(corrida.study, governance=None)
    filas = payload["validation"]["hosmer_lemeshow_groups"]
    assert filas, "el payload de Resultados debe traer los grupos"
    assert set(filas[0]) == set(_COLUMNAS)
    assert {f["partition"] for f in filas} == set(_muestras_con_veredicto(corrida.study))


def test_sdd22_ya_no_dice_bilateral_para_hosmer_lemeshow() -> None:
    """D-HLG-3: el kernel usa la cola superior del χ² (``chi2.sf``); el texto se corrige."""
    texto = (
        Path(__file__).resolve().parents[2] / "docs" / "design" / "22-validation.md"
    ).read_text(encoding="utf-8")
    assert "HL bilateral" not in texto
